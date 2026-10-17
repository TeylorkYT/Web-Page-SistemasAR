import io
from datetime import datetime
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def sanitizar_texto_formula(valor: Any) -> Any:
    """Neutraliza fórmulas maliciosas en hojas de cálculo anteponiendo comilla simple."""
    if isinstance(valor, str):
        if valor.startswith(("=", "+", "-", "@")):
            return f"'{valor}"
    return valor


def _formatear_rango_fechas(desde: str | None, hasta: str | None) -> str:
    if desde and hasta:
        return f"Período: {desde} a {hasta}"
    if desde:
        return f"Período: desde {desde}"
    if hasta:
        return f"Período: hasta {hasta}"
    return "Período: Histórico completo"


def _obtener_metadatos_reporte(tipo_informe: str, datos: Any) -> tuple[str, list[str], list[list[Any]]]:
    """Extrae título, encabezados y filas normalizadas a partir de los datos del informe."""
    if tipo_informe == "por_agente":
        titulo = "Rendimiento por Técnico / Agente"
        encabezados = ["ID Técnico", "Nombre del Técnico", "Tickets Asignados", "Tickets Resueltos", "Tickets Cerrados"]
        filas = []
        for it in datos:
            filas.append([
                it.id_tecnico,
                it.tecnico_nombre,
                it.total,
                it.resueltos,
                it.cerrados,
            ])
        return titulo, encabezados, filas

    if tipo_informe == "por_prioridad":
        titulo = "Distribución de Incidencias por Prioridad"
        encabezados = ["Prioridad", "Total Tickets"]
        filas = [[it.prioridad, it.total] for it in datos]
        return titulo, encabezados, filas

    if tipo_informe == "tiempos_resolucion":
        promedio = getattr(datos, "promedio_horas", 0.0)
        titulo = f"Tiempos de Resolución (Promedio General: {promedio} horas)"
        encabezados = ["ID Ticket", "Tiempo de Resolución (Horas)"]
        filas = [[f"#{t.id_ticket}", t.horas] for t in getattr(datos, "tickets", [])]
        return titulo, encabezados, filas

    if tipo_informe == "estado_general":
        titulo = "Estado General del Sistema"
        encabezados = ["Métrica / Categoría", "Valor"]
        filas = [
            ["Tickets Totales", getattr(datos, "tickets_totales", 0)],
            ["Tickets Abiertos", getattr(datos, "abiertos", 0)],
            ["Tickets Cerrados", getattr(datos, "cerrados", 0)],
            ["Técnicos Activos", getattr(datos, "tecnicos_activos", 0)],
        ]
        for est in getattr(datos, "por_estado", []):
            filas.append([f"Estado: {est.estado}", est.total])
        return titulo, encabezados, filas

    if tipo_informe == "incidencias_area_categoria":
        titulo = "Incidencias por Área y Categoría"
        encabezados = ["Área / Departamento", "Categoría", "Total Incidencias"]
        filas = [[it.area_departamento, it.nombre_categoria, it.total] for it in datos]
        return titulo, encabezados, filas

    # Genérico / Fallback
    titulo = f"Reporte de {tipo_informe.replace('_', ' ').capitalize()}"
    encabezados = ["Registro", "Detalle"]
    filas = [["Total", len(datos) if isinstance(datos, list) else 1]]
    return titulo, encabezados, filas


def generar_excel(
    tipo_informe: str,
    datos: Any,
    desde: str | None = None,
    hasta: str | None = None,
) -> io.BytesIO:
    """Genera un archivo Excel (.xlsx) con formato profesional y neutralización de fórmulas."""
    titulo, encabezados, filas = _obtener_metadatos_reporte(tipo_informe, datos)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = tipo_informe[:31]

    # Estilos
    fuente_titulo = Font(name="Arial", size=14, bold=True, color="0A369D")
    fuente_subtitulo = Font(name="Arial", size=10, italic=True, color="64748B")
    fuente_encabezado = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    relleno_encabezado = PatternFill(start_color="0A369D", end_color="0A369D", fill_type="solid")
    fuente_celda = Font(name="Arial", size=10)
    relleno_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    borde_fino = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # 1. Título y Metadatos
    ws.append([titulo])
    ws.cell(row=1, column=1).font = fuente_titulo
    ws.append([f"{_formatear_rango_fechas(desde, hasta)} · Generado: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
    ws.cell(row=2, column=1).font = fuente_subtitulo
    ws.append([])  # Fila en blanco

    # 2. Encabezados de tabla
    ws.append(encabezados)
    fila_encabezado = 4
    for col_idx in range(1, len(encabezados) + 1):
        celda = ws.cell(row=fila_encabezado, column=col_idx)
        celda.font = fuente_encabezado
        celda.fill = relleno_encabezado
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.border = borde_fino

    # 3. Datos de la tabla
    if not filas:
        ws.append(["Sin datos registrados para el rango seleccionado"] + [""] * (len(encabezados) - 1))
        ws.cell(row=fila_encabezado + 1, column=1).font = Font(name="Arial", size=10, italic=True, color="64748B")
    else:
        for r_idx, fila in enumerate(filas, start=fila_encabezado + 1):
            fila_sanitizada = [sanitizar_texto_formula(celda) for celda in fila]
            ws.append(fila_sanitizada)

            es_impar = (r_idx % 2 == 1)
            for c_idx in range(1, len(fila_sanitizada) + 1):
                celda = ws.cell(row=r_idx, column=c_idx)
                celda.font = fuente_celda
                celda.border = borde_fino
                if es_impar:
                    celda.fill = relleno_zebra
                # Alineación numérica si corresponde
                if isinstance(celda.value, (int, float)):
                    celda.alignment = Alignment(horizontal="right")
                else:
                    celda.alignment = Alignment(horizontal="left")

    # Ajuste automático del ancho de columnas
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generar_pdf(
    tipo_informe: str,
    datos: Any,
    desde: str | None = None,
    hasta: str | None = None,
) -> io.BytesIO:
    """Genera un archivo PDF estructurado y formateado con ReportLab."""
    titulo, encabezados, filas = _obtener_metadatos_reporte(tipo_informe, datos)

    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle(
        "TituloReporte",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#0A369D"),
        spaceAfter=4,
    )
    estilo_meta = ParagraphStyle(
        "MetaReporte",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#64748B"),
        spaceAfter=14,
    )
    estilo_celda = ParagraphStyle(
        "CeldaTexto",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=11,
        textColor=colors.HexColor("#0F172A"),
    )
    estilo_encabezado_celda = ParagraphStyle(
        "EncabezadoTexto",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=11,
        textColor=colors.white,
    )

    elementos = []

    # Encabezado del documento
    elementos.append(Paragraph(titulo, estilo_titulo))
    subtexto = f"{_formatear_rango_fechas(desde, hasta)} · Generado el {datetime.now().strftime('%Y-%m-%d a las %H:%M:%S')}"
    elementos.append(Paragraph(subtexto, estilo_meta))
    elementos.append(Spacer(1, 10))

    # Construcción de la tabla
    tabla_datos = []
    # Encabezados
    tabla_datos.append([Paragraph(enc, estilo_encabezado_celda) for enc in encabezados])

    if not filas:
        tabla_datos.append([Paragraph("Sin datos registrados para el rango seleccionado", estilo_celda)] + [""] * (len(encabezados) - 1))
    else:
        for fila in filas:
            fila_paragrafos = []
            for celda in fila:
                val_str = str(celda if celda is not None else "—")
                fila_paragrafos.append(Paragraph(val_str, estilo_celda))
            tabla_datos.append(fila_paragrafos)

    # Ancho de columnas uniforme ocupando el ancho útil (letter = 612, margen 36*2 = 72 -> útil 540)
    ancho_util = 540
    num_cols = len(encabezados)
    col_width = ancho_util / num_cols

    tabla = Table(tabla_datos, colWidths=[col_width] * num_cols, repeatRows=1)
    tabla.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0A369D")),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ])
    )

    elementos.append(tabla)
    doc.build(elementos)

    output.seek(0)
    return output
